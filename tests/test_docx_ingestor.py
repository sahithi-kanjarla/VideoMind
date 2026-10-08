"""Tests for DOCX ingestor."""
import tempfile
from pathlib import Path
import pytest

from docx import Document

from app.services.ingestion.docx_ingestor import DocxIngestor, _extract_segments


class TestExtractSegments:
    def test_basic_paragraphs(self):
        doc = Document()
        doc.add_paragraph("First paragraph.")
        doc.add_paragraph("Second paragraph.")
        segments = _extract_segments(doc)
        assert len(segments) == 2
        assert segments[0].text == "First paragraph."
        assert segments[1].paragraph_index == 2

    def test_headings_become_sections(self):
        doc = Document()
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("Intro text.")
        doc.add_heading("Methods", level=2)
        doc.add_paragraph("Methods text.")
        segments = _extract_segments(doc)
        assert len(segments) == 2
        assert segments[0].section == "Introduction"
        assert segments[1].section == "Methods"

    def test_empty_paragraphs_skipped(self):
        doc = Document()
        doc.add_paragraph("")
        doc.add_paragraph("Actual text.")
        doc.add_paragraph("")
        segments = _extract_segments(doc)
        assert len(segments) == 1


class TestDocxIngestor:
    def test_file_not_found_raises(self):
        ingestor = DocxIngestor()
        with pytest.raises(ValueError, match="File not found"):
            ingestor.ingest("/nonexistent/file.docx")

    def test_ingest_real_file(self):
        doc = Document()
        doc.add_heading("Test Doc", level=1)
        doc.add_paragraph("This is test content.")

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            doc.save(tmp.name)
            tmp_path = Path(tmp.name)

        try:
            ingestor = DocxIngestor()
            result = ingestor.ingest(tmp_path)
            assert result.source_type == "docx"
            assert len(result.segments) == 1
            assert result.segments[0].section == "Test Doc"
        finally:
            tmp_path.unlink()
